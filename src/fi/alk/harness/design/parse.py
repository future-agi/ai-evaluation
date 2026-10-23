"""Parse one uploaded document into text and tables: `python -m fi.alk.harness.design.parse SRC OUT`.

It runs as a child of the design's control process, with an allow-listed environment and resource
limits (`documents.parse_documents`), because a document is untrusted input: a PDF or a zip-based
office file can be built to exhaust memory or time. It writes `OUT/parts.json` and one file per
part: `<part>.md` for text, `<part>.json` (`{columns, rows}`) for a table.
"""

from __future__ import annotations

import csv
import io
import json
import re
import sys
import zipfile
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

MOST_PAGES = 500
MOST_ROWS = 50_000
MOST_COLUMNS = 200
MOST_CHARS = 2_000_000
# A zip member list or expansion past these is refused before anything is decompressed.
MOST_MEMBERS = 10_000
MOST_UNZIPPED = 200 * 1024 * 1024
MOST_RATIO = 100

_WORD = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


class DocumentError(Exception):
    pass


def kind_of(path: Path) -> str:
    """The document's kind from its signature, not its name alone."""
    head = path.read_bytes()[:8]
    suffix = path.suffix.lower()
    if head.startswith(b"%PDF-"):
        return "pdf"
    if head.startswith(b"PK\x03\x04"):
        if suffix in (".docx", ".xlsx"):
            return suffix[1:]
        raise DocumentError(f"{path.name}: a zip that is not .docx or .xlsx")
    if suffix in (".csv", ".tsv", ".json", ".jsonl", ".txt", ".md"):
        try:
            path.read_bytes().decode("utf-8")
        except UnicodeDecodeError as exc:
            raise DocumentError(f"{path.name}: not valid UTF-8 text") from exc
        return suffix[1:]
    raise DocumentError(f"{path.name}: not a supported document")


def _checked_zip(path: Path) -> zipfile.ZipFile:
    archive = zipfile.ZipFile(path)
    members = archive.infolist()
    unzipped = sum(member.file_size for member in members)
    packed = max(1, sum(member.compress_size for member in members))
    if (
        len(members) > MOST_MEMBERS
        or unzipped > MOST_UNZIPPED
        or unzipped / packed > MOST_RATIO
    ):
        archive.close()
        raise DocumentError(f"{path.name}: expands too far to be read safely")
    return archive


def _cell(value: Any) -> Any:
    if value is None or isinstance(value, (bool, int, float, str)):
        return value
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)


def _table(rows: list[list[Any]]) -> dict[str, Any] | None:
    rows = [row for row in rows if any(cell not in (None, "") for cell in row)]
    if len(rows) < 2:
        return None
    header = [
        str(cell if cell is not None else "").strip() for cell in rows[0][:MOST_COLUMNS]
    ]
    columns = [name or f"column_{index + 1}" for index, name in enumerate(header)]
    body = [
        [_cell(cell) for cell in row[: len(columns)]] for row in rows[1 : MOST_ROWS + 1]
    ]
    return {
        "columns": columns,
        "rows": [row + [None] * (len(columns) - len(row)) for row in body],
        "truncated": len(rows) - 1 > MOST_ROWS,
    }


def _pdf(path: Path) -> list[tuple[str, Any]]:
    from pypdf import PdfReader

    reader = PdfReader(path)
    if len(reader.pages) > MOST_PAGES:
        raise DocumentError(f"{path.name}: more than {MOST_PAGES} pages")
    return [
        (f"page-{number}", page.extract_text() or "")
        for number, page in enumerate(reader.pages, start=1)
    ]


def _docx(path: Path) -> list[tuple[str, Any]]:
    with _checked_zip(path) as archive:
        root = ElementTree.fromstring(archive.read("word/document.xml"))
    body = root.find(f"{_WORD}body")
    if body is None:
        return []
    text: list[str] = []
    parts: list[tuple[str, Any]] = []

    def words(node: ElementTree.Element) -> str:
        return "".join(one.text or "" for one in node.iter(f"{_WORD}t"))

    for node in body:
        if node.tag == f"{_WORD}p":
            text.append(words(node))
        elif node.tag == f"{_WORD}tbl":
            rows = [
                [words(cell) for cell in row.iter(f"{_WORD}tc")]
                for row in node.iter(f"{_WORD}tr")
            ]
            if (table := _table(rows)) is not None:
                parts.append((f"table-{len(parts) + 1}", table))
                text.append(f"[table-{len(parts)}]")
    return [("text", "\n".join(text).strip()), *parts]


def _xlsx(path: Path) -> list[tuple[str, Any]]:
    from openpyxl import load_workbook

    _checked_zip(path).close()
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        parts = []
        for sheet in workbook.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                rows.append(list(row))
                if len(rows) > MOST_ROWS + 1:
                    break
            if (table := _table(rows)) is not None:
                parts.append((sheet.title, table))
        return parts
    finally:
        workbook.close()


def _delimited(path: Path, delimiter: str) -> list[tuple[str, Any]]:
    rows = list(
        csv.reader(io.StringIO(path.read_text(encoding="utf-8")), delimiter=delimiter)
    )
    table = _table(rows)
    return [("table", table)] if table is not None else []


def _json(path: Path, lines: bool) -> list[tuple[str, Any]]:
    text = path.read_text(encoding="utf-8")
    try:
        value = (
            [json.loads(line) for line in text.splitlines() if line.strip()]
            if lines
            else json.loads(text)
        )
    except ValueError as exc:
        raise DocumentError(f"{path.name}: not valid JSON") from exc
    if (
        isinstance(value, list)
        and value
        and all(isinstance(one, dict) for one in value)
    ):
        columns = list(dict.fromkeys(key for one in value for key in one))
        table = _table([columns, *([one.get(key) for key in columns] for one in value)])
        if table is not None:
            return [("table", table)]
    return [("text", json.dumps(value, indent=2, ensure_ascii=False))]


def parse(source: Path) -> tuple[str, list[tuple[str, Any]]]:
    kind = kind_of(source)
    parts = {
        "pdf": _pdf,
        "docx": _docx,
        "xlsx": _xlsx,
        "csv": lambda path: _delimited(path, ","),
        "tsv": lambda path: _delimited(path, "\t"),
        "json": lambda path: _json(path, lines=False),
        "jsonl": lambda path: _json(path, lines=True),
        "txt": lambda path: [("text", path.read_text(encoding="utf-8"))],
        "md": lambda path: [("text", path.read_text(encoding="utf-8"))],
    }[kind](source)
    return kind, parts


def part_name(raw: str) -> str:
    return re.sub(r"[^A-Za-z0-9_.-]+", "-", raw).strip("-.") or "part"


def main(argv: list[str]) -> int:
    source, out = Path(argv[0]), Path(argv[1])
    try:
        kind, parts = parse(source)
    except DocumentError as exc:
        print(str(exc), file=sys.stderr)
        return 3
    out.mkdir(parents=True, exist_ok=True)
    listed, total = [], 0
    for raw, content in parts:
        name = part_name(raw)
        if isinstance(content, str):
            total += len(content)
            if total > MOST_CHARS:
                print(
                    f"{source.name}: more than {MOST_CHARS} characters of text",
                    file=sys.stderr,
                )
                return 3
            (out / f"{name}.md").write_text(content, encoding="utf-8")
            listed.append({"name": name, "kind": "text", "chars": len(content)})
        else:
            (out / f"{name}.json").write_text(
                json.dumps(content, ensure_ascii=False), encoding="utf-8"
            )
            listed.append(
                {
                    "name": name,
                    "kind": "table",
                    "columns": content["columns"],
                    "rows": len(content["rows"]),
                    "truncated": content["truncated"],
                }
            )
    (out / "parts.json").write_text(
        json.dumps({"kind": kind, "parts": listed}, indent=2), encoding="utf-8"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
